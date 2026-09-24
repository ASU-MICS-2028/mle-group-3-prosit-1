"""Build the Prosit 1 panel deck on the official Ashesi "Presentation Red" template.

Run from the repo root:  .venv/bin/python scripts/build_ashesi_deck.py
Every number on a slide comes from a notebook cell or scripts/verify_claims.py; each
slide's speaker notes say which. Writes reports/ITN_Allocation_Ashesi.pptx.
"""

import sys
import tempfile
import warnings
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_TICK_LABEL_POSITION,
    XL_TICK_MARK,
)
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

warnings.filterwarnings("ignore")
HERE = Path(__file__).resolve().parent
REPO = HERE.parent
TEMPLATE = HERE / "ashesi_presentation_red.pptx"
OUT = REPO / "reports" / "ITN_Allocation_Ashesi.pptx"

RED, GOLD, INK, MUTED = "AB3D3F", "FEBA5A", "262626", "595959"
CARD, GRID, GREY, WHITE = "F3F3F3", "D9D9D9", "A6A6A6", "FFFFFF"
FONT = "Candara"  # template body font; titles keep the layouts' Poppins
TITLE, IMAGE_TEXT, TEXT_HEAVY, TABLE, CLOSING = 0, 1, 3, 7, 8  # template layout indices
W = 12.09  # content width from the 0.62" margin; the wordmark sits below y = 6.5"


# ---------------------------------------------------------------- text helpers
def rgb(h):
    return RGBColor.from_string(h)


def P(*runs, bullet=False, after=None, before=None, align=None, s=None, c=None):
    """One paragraph: runs are str or (str, style) with style keys b, c, s, i."""
    return dict(
        runs=runs, bullet=bullet, after=after, before=before, align=align, s=s, c=c
    )


def B(t, **st):
    return (t, {"b": True, **st})


def bulletize(p, color=RED, indent=0.24):
    ppr = p._p.get_or_add_pPr()
    ppr.set("marL", str(Inches(indent)))
    ppr.set("indent", str(-Inches(indent)))
    for tag in ("a:buNone", "a:buClr", "a:buFont", "a:buChar", "a:buAutoNum"):
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    etree.SubElement(etree.SubElement(ppr, qn("a:buClr")), qn("a:srgbClr")).set(
        "val", color
    )
    etree.SubElement(ppr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(ppr, qn("a:buChar")).set("char", "•")


def fill(tf, paras, size=14, color=INK, align=PP_ALIGN.LEFT, space=6):
    tf.clear()
    for i, para in enumerate(paras):
        para = para if isinstance(para, dict) else P(para)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para["align"] or align
        p.space_after = Pt(space if para["after"] is None else para["after"])
        if para["before"] is not None:
            p.space_before = Pt(para["before"])
        for run in para["runs"]:
            t, st = (run, {}) if isinstance(run, str) else run
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(st.get("s", para["s"] or size))
            f.bold = st.get("b", False)
            f.italic = st.get("i", False)
            f.color.rgb = rgb(st.get("c", para["c"] or color))
        if para["bullet"]:
            bulletize(p)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill(tf, paras, **kw)
    return tb


def card(slide, x, y, w, h, paras, bg=CARD, pad=0.22, anchor=MSO_ANCHOR.TOP, **kw):
    s = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    s.adjustments[0] = 0.05
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(bg)
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(pad * 0.8)
    if paras:
        fill(tf, paras, **kw)
    return s


def title(slide, idx, t, size=None, width=None):
    ph = slide.placeholders[idx]
    if width:  # override all four so the inherited position is kept
        ph.left, ph.top, ph.width, ph.height = ph.left, ph.top, Inches(width), ph.height
    ph.text_frame.text = t
    if size:
        ph.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    return ph


def drop(slide, *idxs):
    for i in idxs:
        el = slide.placeholders[i]._element
        el.getparent().remove(el)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def picture(slide, path, x, y, w, h):
    iw, ih = Image.open(path).size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    return slide.shapes.add_picture(
        str(path),
        Inches(x + (w - pw) / 2),
        Inches(y + (h - ph) / 2),
        Inches(pw),
        Inches(ph),
    )


# ---------------------------------------------------------------- charts and tables
def bar_chart(
    slide, x, y, w, h, cats, vals, colors, fmt="#,##0", horizontal=False, axis=True
):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("value", vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    ch = slide.shapes.add_chart(
        kind, Inches(x), Inches(y), Inches(w), Inches(h), cd
    ).chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = FONT
    ch.font.size = Pt(12)
    ch.font.color.rgb = rgb(INK)
    plot = ch.plots[0]
    plot.gap_width = 60
    series = plot.series[0]
    series.invert_if_negative = False
    for pt, c in zip(series.points, colors):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
    for dpt in series._element.findall(
        qn("c:dPt")
    ):  # PowerPoint inverts negative points without this
        flag = etree.Element(qn("c:invertIfNegative"))
        flag.set("val", "0")
        dpt.find(qn("c:idx")).addnext(flag)
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold = Pt(13), True
    va, ca = ch.value_axis, ch.category_axis
    va.has_major_gridlines = axis
    if axis:
        va.major_gridlines.format.line.color.rgb = rgb(GRID)
    va.format.line.fill.background()
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.tick_labels.font.size = Pt(11)
    va.tick_labels.font.color.rgb = rgb(MUTED)
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = fmt, False
    va.visible = axis
    ca.format.line.color.rgb = rgb(GREY)
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.tick_labels.font.size = Pt(12)
    if horizontal:
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    return ch


def table(slide, x, y, col_w, rows, row_h, size=12, right_cols=(), bold_rows=()):
    tbl = slide.shapes.add_table(
        len(rows),
        len(col_w),
        Inches(x),
        Inches(y),
        Inches(sum(col_w)),
        Inches(sum(row_h)),
    ).table
    tbl.first_row, tbl.horz_banding = True, False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h[i])
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(
                RED if i == 0 else (WHITE if i % 2 else CARD)
            )
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            bold = i == 0 or j == 0 or i in bold_rows
            align = PP_ALIGN.RIGHT if j in right_cols else PP_ALIGN.LEFT
            fill(
                cell.text_frame,
                [P((val, {"b": bold}))],
                size=size + (1 if i == 0 else 0),
                color=WHITE if i == 0 else INK,
                align=align,
                space=0,
            )
    return tbl


# ---------------------------------------------------------------- figures
def crop_title(src, dst, gap=12):
    """Drop a matplotlib figure's baked-in title: keep from the second ink band down."""
    im = Image.open(src).convert("RGB")
    ink = (np.asarray(im.convert("L")) < 200).any(axis=1)
    bands, start, white = [], None, 0
    for i, dark in enumerate(ink):
        if dark:
            start, white = (i if start is None else start), 0
        elif start is not None:
            white += 1
            if white >= gap:
                bands.append(start)
                start, white = None, 0
    if start is not None:
        bands.append(start)
    top = max(bands[1] - 12, 0) if len(bands) > 1 else 0
    im.crop((0, top, im.width, im.height)).save(dst)
    return dst


def north_maps(tmp):
    """Draw the allocation, burden and uncertainty maps with the team's own code, zoomed north."""
    import matplotlib.pyplot as plt

    sys.path.insert(0, str(REPO))
    from src import io as pio, models, viz

    d = pio.load_district_cases()
    off = np.log(d["mean_population"].clip(lower=1)).to_numpy()
    nb = models.fit_negative_binomial_mle(
        "positive_cases ~ net_coverage_pct", d, offset=off
    )
    alloc = models.compute_allocation(d, nb, total_nets=50000, offset=off)
    assert alloc["equitable_allocation"].sum() == 50000
    alloc["cases_per_person"] = alloc["positive_cases"] / alloc["mean_population"]
    # how far the upper 95% bound sits above the estimate; one value per region
    alloc["pct_above"] = (alloc["pred_ci_upper"] / alloc["predicted_cases"] - 1) * 100
    by_region = alloc.groupby("region_code")["pct_above"]
    assert (by_region.max() - by_region.min()).max() < 1e-6
    assert sorted(by_region.first().round()) == [16, 20, 34]  # the slide text
    b = REPO / "data" / "ghana_boundaries"
    viz.set_theme()
    out = {}
    for key, col, label in [
        ("alloc", "equitable_allocation", "Nets allocated under the proposed rule"),
        ("burden", "cases_per_person", "Confirmed cases per person, 2014-17"),
        ("uncertainty", "pct_above", "Upper 95% bound, % above the estimate"),
    ]:
        fig, ax = viz.plot_district_choropleth(
            b / "gha_admin0.geojson",
            b / "gha_admin1.geojson",
            b / "gha_admin2.geojson",
            alloc,
            value_col=col,
            title="",
        )
        ax.set_xlim(-3.3, 0.75)
        ax.set_ylim(7.9, 11.25)
        cbar = fig.axes[-1]
        cbar.set_xlabel(label, fontsize=13)
        cbar.tick_params(labelsize=11)
        out[key] = tmp / f"{key}_map.png"
        fig.savefig(out[key], dpi=220, bbox_inches="tight", pad_inches=0.05)
        plt.close(fig)
    return out


# ---------------------------------------------------------------- the deck
def build():
    tmp = Path(tempfile.mkdtemp(prefix="prosit1_deck_"))
    fig = {
        n: crop_title(REPO / "figures" / f"{n}.png", tmp / f"{n}.png")
        for n in (
            "a3_poisson_vs_nb",
            "a4_bootstrap_ci_comparison",
            "b1_data_availability",
            "b1_district_case_rate",
        )
    }
    fig.update(north_maps(tmp))

    prs = Presentation(TEMPLATE)
    n_template = len(prs.slides)
    add = lambda layout: prs.slides.add_slide(prs.slide_layouts[layout])  # noqa: E731

    # 1. Title
    s = add(TITLE)
    title(s, 0, "Where should 50,000 bed nets go?", size=60)
    fill(
        s.placeholders[1].text_frame,
        [
            P(
                "Allocating bed nets across northern Ghana, with honest uncertainty",
                s=22,
                c=GOLD,
                after=10,
            ),
            P(
                "Advisory briefing for the National Malaria Elimination Programme",
                s=18,
                c=WHITE,
                after=4,
            ),
            P(
                "ICS553 Machine Learning Essentials · Group 3 · Ashesi University · September 2026",
                s=16,
                c=WHITE,
            ),
        ],
        align=PP_ALIGN.CENTER,
    )
    notes(
        s,
        "Say: We were asked where a limited delivery of 50,000 insecticide-treated nets should go "
        "across northern Ghana, and how confident we can be in the district numbers behind that choice. "
        "We start with our recommendation, then show the evidence and its limits.\n"
        "Add the presenters' names to this slide before the panel.",
    )

    # 2. Recommendation
    s = add(TEXT_HEAVY)
    title(s, 11, "Our recommendation", width=W)
    drop(s, 12, 13)
    text(
        s,
        0.62,
        1.6,
        W,
        0.8,
        [
            P(
                "For a 50,000-net consignment to the three northern regions: ",
                B("allocate within each region in proportion to population"),
                ", with regional totals from our negative binomial model.",
            )
        ],
        size=18,
    )
    for i, (num, lab, sub) in enumerate(
        [
            (
                "24,196",
                "nets for Northern Region",
                "26 districts · 8.5 nets per 1,000 people",
            ),
            (
                "18,631",
                "nets for Upper East",
                "13 districts · 16.0 nets per 1,000 people",
            ),
            (
                "7,173",
                "nets for Upper West",
                "11 districts · 9.2 nets per 1,000 people",
            ),
        ]
    ):
        card(
            s,
            0.62 + i * 4.13,
            2.65,
            3.83,
            2.25,
            [
                P(B(num, c=RED, s=46), after=2),
                P(B(lab, s=17), after=4),
                P(sub, s=14, c=MUTED),
            ],
            anchor=MSO_ANCHOR.MIDDLE,
        )
    text(
        s,
        0.62,
        5.1,
        W,
        1.3,
        [
            P(
                B("Why population within a region: ", c=RED),
                "net coverage is measured once per region, so our data cannot tell two districts "
                "in the same region apart.",
                after=6,
            ),
            P(
                B("How sure we are: ", c=RED),
                "the shares within each region are robust. The regional totals are our base case: "
                "a better-fitting model moves 7,720 nets between regions (slide 13).",
            ),
        ],
        size=15,
    )
    notes(
        s,
        "Say: Within each region the rule allocates in proportion to population, because coverage is "
        "measured once per region. The regional totals come from our negative binomial model and are our "
        "base case; be upfront that the regional split depends on the model.\n"
        "Sources: notebooks/04_allocation.ipynb, cell alloc_cd07 (regional totals), alloc_cd11 "
        "(nets per 1,000 people) and alloc_cd13 (7,720 nets moved under the region-effects model).",
    )

    # 3. Averages hide local risk
    s = add(TEXT_HEAVY)
    title(s, 11, "Averages hide local risk", width=W)
    drop(s, 12, 13)
    card(
        s,
        0.62,
        1.7,
        5.9,
        3.0,
        [
            P(B("2%", c=RED, s=66), after=4),
            P(B("Greater Accra malaria prevalence, 2022 DHS", s=18), after=4),
            P("Children 6 to 59 months, microscopy", s=14, c=MUTED),
        ],
        pad=0.3,
    )
    card(
        s,
        6.81,
        1.7,
        5.9,
        3.0,
        [
            P(B("0 to 49%", c=RED, s=66), after=4),
            P(
                B("Predicted prevalence across all 29 Greater Accra districts", s=18),
                after=4,
            ),
            P(
                "2020 household survey and geostatistical model, children 6 months to 10 years, rapid tests",
                s=14,
                c=MUTED,
            ),
        ],
        pad=0.3,
    )
    text(
        s,
        0.62,
        5.15,
        W,
        1.1,
        [
            P(
                "Different ages and tests, so this is not like for like. The point is the spread: one regional figure "
                "can hide districts at very high risk. That is why we wanted district-level targeting; the next "
                "slide shows how far our data allow it."
            )
        ],
        size=16,
    )
    notes(
        s,
        "Say: This is why district detail matters. Be careful with the wording: the 49% is a model "
        "prediction for children aged 6 months to 10 years, and its highest values fall in forest areas, "
        "not slums. The 2% is the 2022 DHS microscopy figure for children under five (3.4% by rapid test).\n"
        "Sources: Oppong et al., Malaria Journal, doi:10.1186/s12936-025-05724-9 "
        "(https://pmc.ncbi.nlm.nih.gov/articles/PMC12829291/); Ghana DHS 2022 final report FR387, "
        "Table 12.15.",
    )

    # 4. Data and where it is thin
    s = add(TEXT_HEAVY)
    title(s, 11, "Our data, and where it is thin", width=W)
    drop(s, 12, 13)
    picture(s, fig["b1_data_availability"], 0.62, 1.55, 3.55, 3.95)
    picture(s, fig["b1_district_case_rate"], 4.42, 1.55, 3.55, 3.95)
    text(
        s,
        0.62,
        5.6,
        3.55,
        0.8,
        [
            P(
                B("Household survey, DHS 2022. "),
                "17,933 households in 618 clusters, all 16 regions; no district identifier.",
            )
        ],
        size=12,
    )
    text(
        s,
        4.42,
        5.6,
        3.55,
        0.8,
        [
            P(
                B("Routine surveillance, 2014-17. "),
                "Cases for 50 northern districts only; 207 of today's 260 districts have none in our data.",
            )
        ],
        size=12,
    )
    card(
        s,
        8.35,
        1.55,
        4.36,
        4.75,
        [
            P(B("What this means", c=RED, s=18), after=8),
            P(
                B("Weights matter. "),
                "Unweighted, 71.0% of households own a net; weighted, 66.8%.",
                bullet=True,
                after=8,
            ),
            P(
                B("Different years. "),
                "Net coverage is from 2022; cases are from 2014-17.",
                bullet=True,
                after=8,
            ),
            P(
                B("Coarse coverage. "),
                "Coverage is known only by region: three values for 50 districts.",
                bullet=True,
                after=8,
            ),
            P(
                B("No district survey data. "),
                "The survey can rank regions, not districts.",
                bullet=True,
                after=12,
            ),
            P(
                B("So: ", c=RED),
                "our rule, expected cases × the share of households without a net, can separate "
                "regions but not districts within a region.",
            ),
        ],
        size=15,
    )
    notes(
        s,
        "Say: We have two very different sources. The survey is representative but only resolves to "
        "region; the surveillance has district detail but only for the north.\n"
        "Sources: 17,933 households and 618 clusters: data dictionary and notebooks/02 cell cd03. "
        "71.0% vs 66.8% (claim C-01): notebooks/01_eda.ipynb cell eda_cd05. 50 districts on 53 polygons, "
        "207 of 260 districts without data: scripts/verify_claims.py (map join, same logic as src/viz.py). "
        "Maps: figures/b1_data_availability.png and figures/b1_district_case_rate.png from notebook 01, "
        "with their titles trimmed.",
    )

    # 5. Over-dispersion
    s = add(TEXT_HEAVY)
    title(s, 11, "Case counts vary far more than Poisson allows", size=30, width=W)
    drop(s, 12, 13)
    for i, (num, desc) in enumerate(
        [
            (
                "77,183",
                "variance-to-mean ratio of district case counts; a Poisson model requires 1",
            ),
            (
                "442 vs 122,882",
                "spread a Poisson implies (standard deviation) against what we observe: "
                "278 times too narrow",
            ),
            (
                "54,147",
                "dispersion left after adjusting for population (Pearson chi-square per degree of freedom)",
            ),
        ]
    ):
        text(
            s,
            0.62,
            1.6 + i * 1.12,
            5.2,
            1.05,
            [P(B(num, c=RED, s=30), after=0), P(desc, s=13)],
        )
    card(
        s,
        0.62,
        5.05,
        5.2,
        1.2,
        [
            P(
                B("Working model: negative binomial. ", c=RED),
                "Dispersion α = 0.268 with a population offset; it beats Poisson by 2.39 million AIC points.",
            )
        ],
        size=14,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    picture(s, fig["a3_poisson_vs_nb"], 6.2, 1.55, 6.5, 4.15)
    text(
        s,
        6.2,
        5.8,
        6.5,
        0.5,
        [
            P(
                "50 districts, 2014-17. Curves fitted without covariates, so the negative "
                "binomial here uses α = 0.395."
            )
        ],
        size=11,
        color=MUTED,
    )
    notes(
        s,
        "Say: Under a Poisson model the variance equals the mean. Here the variance is 77,183 times the "
        "mean, and even after adjusting for population the dispersion ratio is 54,147. The negative "
        "binomial adds one dispersion parameter and fits far better. The figure's curves use no "
        "covariates (alpha 0.395); the model we use has a population offset (alpha 0.268).\n"
        "Sources: notebooks/02_distributions.ipynb cells cd09 (77,183), cd15 (442, 122,882, 278 times), "
        "cd18 (0.395), cd25 and cd26 (54,147; alpha 0.2677; AIC gap 2,393,669). "
        "Figure: figures/a3_poisson_vs_nb.png, title trimmed.",
    )

    # 6. Survey intervals
    s = add(TEXT_HEAVY)
    title(s, 11, "Survey intervals must respect clustering", width=W)
    drop(s, 12, 13)
    text(
        s,
        0.62,
        1.6,
        5.6,
        4.75,
        [
            P(
                "Households in one cluster share a village, water sources and the last net campaign, so they are "
                "not independent.",
                after=12,
            ),
            P(B("15.2 vs 7.2 points", c=RED, s=22), after=2),
            P(
                "Width of the 95% interval for rural Northern net ownership (76.0%): cluster bootstrap against a "
                "naive household bootstrap. The honest interval is 2.1 times wider.",
                after=12,
            ),
            P(B("Design effect about 4.5", c=RED, s=22), after=2),
            P(
                "The survey holds the information of a simple random sample about a quarter its size "
                "(4.5 to 5.4 across seeds; 4.2 by a design-based formula).",
                after=12,
            ),
            P(B("0.09 points", c=RED, s=22), after=2),
            P(
                "Our weighted estimate for Northern Region against the published DHS figure (67.69% vs 67.60%). "
                "This checks our weighting, not our interval: the published ITN intervals are empty."
            ),
        ],
        size=14,
    )
    picture(s, fig["a4_bootstrap_ci_comparison"], 6.5, 1.9, 6.2, 3.1)
    text(
        s,
        6.5,
        5.2,
        6.2,
        1.0,
        [
            P(
                "Rural Northern Region: 582 households in 20 clusters, 2,000 replicates, "
                "seed 42. Chosen for its link to our districts; it ranks 10th of 16 "
                "regions by rural clusters, so it is not the most under-sampled."
            )
        ],
        size=11,
        color=MUTED,
    )
    notes(
        s,
        "Say: The survey samples clusters, then households. Resampling households alone pretends each "
        "is independent and gives an interval half as wide as it should be.\n"
        "Sources: notebooks/02_distributions.ipynb cells cd31 (582 households, 20 clusters), cd33 and "
        "cd34 (15.20 vs 7.17, 2.12 times, design effect 4.50), cd37 (published ITN intervals empty, "
        "0 of 76), cd39 (67.69% vs 67.6%). Seed spread 4.46 to 5.39 and design-based 4.18: "
        "scripts/verify_claims.py section A. Rank 10th of 16: scripts/verify_claims.py section D. "
        "Figure: figures/a4_bootstrap_ci_comparison.png, title trimmed.",
    )

    # 7. Leakage audit
    s = add(TABLE)
    title(s, 11, "Leakage audit: what held up, and what did not", size=30, width=11.01)
    drop(s, 10)
    table(
        s,
        1.16,
        1.6,
        [1.75, 2.85, 2.85, 3.56],
        [
            ["Leak", "Broken version", "Disciplined version", "What we found"],
            [
                "Preprocessing",
                "Scaler fitted on all 50 districts before splitting",
                "Scaler fitted inside a Pipeline, on training districts only",
                "Real mechanism, negligible here: 212 cases at seed 42, 44 on average over 500 seeds",
            ],
            [
                "Target encoding",
                "Each district encoded with its own case count",
                "Only features that do not use the outcome",
                "Test R² of 1.00 against 0.26: the feature is the answer",
            ],
            [
                "Spatial",
                "Random split, or random split stratified by region",
                "Whole region held out",
                "Random and stratified splits agree (about 90,000); a held-out region gives 365,128",
            ],
            [
                "Temporal",
                "2022 coverage used to explain 2014-17 cases",
                "Documented; not read causally",
                "The coverage coefficient reflects region, not the effect of nets",
            ],
        ],
        [0.45, 1.0, 0.85, 1.0, 0.85],
        size=12,
    )
    notes(
        s,
        "Say: We built each leak deliberately and compared it with the disciplined version. The "
        "spatial one needed a different test from the one we first used: comparing a random split with a "
        "stratified split at a single seed compares two different random test sets, and across seeds the "
        "difference disappears.\n"
        "Sources: notebooks/03_pipeline_leakage.ipynb cells leak_cd04 (87,144 vs 87,356), leak_cd06 "
        "(R² 1.0000 vs 0.2618), leak_cd08 (one seed), leak_cd12 (500 seeds) and leak_cd14 "
        "(whole region held out); the table itself is leak_cd10.",
    )

    # 8. Spatial leakage chart
    s = add(TEXT_HEAVY)
    title(s, 11, "Only a held-out region tests new geography", width=W)
    drop(s, 12, 13)
    text(
        s,
        0.62,
        1.7,
        5.2,
        4.5,
        [
            P(
                "A split stratified by region still puts each test district's neighbours in the training set.",
                bullet=True,
                after=12,
            ),
            P(
                "Across 500 random seeds, random and region-stratified splits give the same test error.",
                bullet=True,
                after=12,
            ),
            P(
                "Train on two regions, predict the third, and the error is four times larger.",
                bullet=True,
                after=12,
            ),
            P(
                "So the model cannot yet rank districts in a region it has not seen, such as the 207 districts "
                "outside our data.",
                bullet=True,
            ),
        ],
        size=16,
    )
    bar_chart(
        s,
        6.1,
        1.5,
        6.6,
        4.25,
        ["Random split", "Stratified by region", "Region held out"],
        [91413, 90169, 365128],
        [GREY, GREY, RED],
    )
    text(
        s,
        6.1,
        5.8,
        6.6,
        0.55,
        [
            P(
                "Test RMSE in confirmed cases, Ridge model from notebook 03. Random splits: "
                "median of 500 seeds. Held out: Northern 272,529, Upper East 589,777, "
                "Upper West 139,143."
            )
        ],
        size=11,
        color=MUTED,
    )
    notes(
        s,
        "Say: Stratifying by region keeps every region in both training and test, so neighbours still "
        "leak. Only holding a whole region out tests what the model does somewhere new, and the error "
        "roughly quadruples.\n"
        "Sources: notebooks/03_pipeline_leakage.ipynb, cells leak_cd12 (500 seeds) and leak_cd14 "
        "(whole region held out, via io.split_data(..., hold_out=region)). Literature: Ploton et al. 2020, Nature Communications 11:4540, "
        "doi:10.1038/s41467-020-18321-y; Kattenborn et al. 2022, ISPRS Open Journal of Photogrammetry and "
        "Remote Sensing, doi:10.1016/j.ophoto.2022.100018 (up to 28% overestimation of F1-score).",
    )

    # 9. What the rule does
    s = add(TEXT_HEAVY)
    title(s, 11, "What our rule does with our data", width=W)
    drop(s, 12, 13)
    text(
        s,
        0.62,
        1.55,
        6.0,
        4.8,
        [
            P(B("The rule", c=RED, s=18), after=4),
            P(
                "Weight = upper 95% bound of expected cases (negative binomial with a population offset) × "
                "(1 - net coverage). Hamilton's largest-remainder method turns the weights into exactly 50,000 "
                "whole nets.",
                after=14,
            ),
            P(B("What it does in practice", c=RED, s=18), after=4),
            P(
                "Coverage is one number per region, so expected cases are population times a regional rate. "
                "Within a region, nets follow population exactly.",
                bullet=True,
                after=8,
            ),
            P(
                "The upper bound is a confidence interval for the mean, not a prediction interval. It adds one "
                "multiplier per region: 1.20, 1.34 and 1.16.",
                bullet=True,
                after=8,
            ),
            P(
                "The coverage coefficient is positive (+0.082 per point): regions with more nets had more malaria. "
                "So higher coverage earns more nets per person.",
                bullet=True,
            ),
        ],
        size=14,
    )
    bar_chart(
        s,
        7.0,
        1.55,
        5.7,
        4.2,
        ["Northern (67.7%)", "Upper West (69.8%)", "Upper East (79.6%)"],
        [8.51, 9.17, 16.02],
        [RED, RED, RED],
        fmt="0.0",
    )
    text(
        s,
        7.0,
        5.8,
        5.7,
        0.55,
        [
            P(
                "Nets per 1,000 people by region, with the share of households owning a net (DHS 2022). "
                "Upper East has the highest coverage and the most nets per person: the coefficient "
                "reflects region, not the effect of nets, so the regional split is uncertain (slide 13)."
            )
        ],
        size=11,
        color=MUTED,
    )
    notes(
        s,
        "Say: Because coverage is a single regional number, the model can only tell regions apart, so "
        "within a region every district gets the same nets per person. The coverage coefficient is "
        "positive because nets went where malaria was worst, and coverage is from 2022 while cases are "
        "from 2014-17, so it cannot be read as a protective effect.\n"
        "Sources: notebooks/04_allocation.ipynb cell alloc_cd11 (coefficient +0.0823, p = 1.7e-07; "
        "upper-to-expected ratios 1.198, 1.342, 1.162; nets per 1,000 people). Coverage values: "
        "notebooks/02 cell cd06 and scripts/verify_claims.py section D.",
    )

    # 10. Burden next to uncertainty: the two-map slide the course roadmap asks for
    s = add(TEXT_HEAVY)
    title(s, 11, "Where the need looks highest, and how sure we are", size=30, width=W)
    drop(s, 12, 13)
    for x, key, head, body in [
        (
            0.62,
            "burden",
            "Reported burden. ",
            "Confirmed cases per person, 2014-17: highest in Upper East and Upper West, "
            "lowest across Northern Region. Grey areas have no case data of their own.",
        ),
        (
            6.81,
            "uncertainty",
            "How sure we are. ",
            "Our upper bound sits 20% above the estimate in Northern, 34% in Upper East and "
            "16% in Upper West: one number per region, because coverage is.",
        ),
    ]:
        picture(s, fig[key], x, 1.45, 5.9, 3.8)
        text(s, x, 5.35, 5.9, 0.6, [P(B(head), body)], size=12)
    text(
        s,
        0.62,
        6.02,
        W,
        0.4,
        [
            P(
                B("Takeaway: ", c=RED),
                "the rule uses the upper bound, so it leans towards the region we are least sure about.",
            )
        ],
        size=14,
    )
    notes(
        s,
        "Say: The left map shows reported burden: confirmed cases per person over 2014-17. It is highest "
        "in Upper East and Upper West and lowest across Northern Region, where the test data point to "
        "under-testing. The right map shows how sure we are: the upper end of our 95% interval sits 20% "
        "above the estimate in Northern, 34% in Upper East and 16% in Upper West. It is one number per "
        "region because coverage is. Our rule uses that upper bound, so it leans towards the region we "
        "are least sure about.\n"
        "Sources: cases per person and the upper-to-expected ratios 1.198, 1.342 and 1.162: "
        "notebooks/04_allocation.ipynb cell alloc_cd11. Both maps are drawn with "
        "src/viz.plot_district_choropleth from models.compute_allocation, zoomed to the north like the "
        "allocation map. The two-map layout (burden next to uncertainty) follows the course roadmap.",
    )

    # 11. Proposed allocation map
    s = add(IMAGE_TEXT)
    title(s, 11, "The proposed allocation", size=30, width=6.2)
    drop(s, 12)
    fill(
        s.placeholders[10].text_frame,
        [
            P(B("Regional totals", c=RED, s=18), after=4),
            P("Northern 24,196 · Upper East 18,631 · Upper West 7,173", after=14),
            P(B("Largest allocations", c=RED, s=18), after=4),
            P(
                "Bolgatanga 2,350 · Garu-Tempane 2,309 · Tamale 2,175 · Kasena-Nankana 1,954 "
                "· Bawku 1,750",
                after=14,
            ),
            P(B("Smallest", c=RED, s=18), after=4),
            P(
                "Daffiama-Bussie-Issa, 335. Every district receives nets, and the total is exactly 50,000."
            ),
        ],
        size=16,
    )
    card(s, 7.5, 0.7, 5.45, 5.75, None, bg=WHITE)
    picture(s, fig["alloc"], 7.65, 0.85, 5.15, 5.45)
    notes(
        s,
        "Say: This is the allocation our current rule produces.\n"
        "Sources: the 50-district schedule in reports/allocation.md reproduces exactly from "
        "notebooks/04_allocation.ipynb (cells alloc_cd07 and alloc_cd09; scripts/verify_claims.py "
        "section C checks all 50 rows). Map drawn with src/viz.plot_district_choropleth from the same "
        "allocation, zoomed to the north. The three pre-2018 districts later split (Garu-Tempane, "
        "Savelugu-Nanton, Bunkpurugu-Yunyoo) show their total on both successor areas.",
    )

    # 12. Gains and losses
    s = add(TEXT_HEAVY)
    title(s, 11, "Who gains, who loses, and why", width=W)
    drop(s, 12, 13)
    text(
        s,
        0.62,
        1.55,
        5.5,
        4.85,
        [
            P(
                "Our comparator gives nets in proportion to reported cases. It is not Ghana's practice: national "
                "campaigns allocate by population, about one net per two people.",
                bullet=True,
                after=9,
            ),
            P(
                "Tamale gains most (+1,692): the largest population, and the lowest reported cases per person of "
                "all 50 districts.",
                bullet=True,
                after=9,
            ),
            P(
                "Within each region, the more cases per person a district reported, the more nets per person it "
                "gives up. Nabdam and Bole, the highest in their regions, lose 761 and 376.",
                bullet=True,
                after=9,
            ),
            P(
                "We treat everyone in a region as equally at risk. That is a value judgement, and we make it openly.",
                bullet=True,
                after=9,
            ),
            P(
                "Referral bias at Wa and Bolgatanga is a hypothesis, not a finding: Tamale hosts the north's only "
                "tertiary hospital yet reports the lowest rate.",
                bullet=True,
            ),
        ],
        size=14,
    )
    names = [
        "Wa",
        "Nabdam",
        "Lambussie-Karni",
        "Bolgatanga",
        "Sissala East",
        "Gushiegu",
        "West Mamprusi",
        "East Gonja",
        "Sagnarigu",
        "Tamale",
    ]
    deltas = [-1516, -761, -707, -665, -634, 690, 791, 878, 1012, 1692]
    bar_chart(
        s,
        6.35,
        1.45,
        6.4,
        4.4,
        names,
        deltas,
        [GREY if v < 0 else RED for v in deltas],
        fmt="+#,##0;-#,##0",
        horizontal=True,
        axis=False,
    )
    text(
        s,
        6.35,
        5.9,
        6.4,
        0.45,
        [
            P(
                "Change in nets against case-proportional allocation: the five largest "
                "gains and losses."
            )
        ],
        size=11,
        color=MUTED,
    )
    notes(
        s,
        "Say: Against a case-proportional comparator, nets move from districts with high reported cases "
        "per person to populous districts. That follows from treating everyone in a region as equally at "
        "risk. Referral bias may exist, but our data do not show it: Bolgatanga ranks only 4th of 13 in "
        "Upper East on cases per person, and Nabdam, often named as a sending district, ranks 1st.\n"
        "Sources: gains and losses: notebooks/04_allocation.ipynb alloc_cd09; cases-per-person ranks: "
        "alloc_cd11 (Tamale lowest of 50, Nabdam highest of 50, Bole highest in its region). Ghana's ITN strategy: PMI Ghana Malaria Operational Plan FY2017. Tamale "
        "Teaching Hospital as the only tertiary provider for the three northern regions: "
        "https://tth.gov.gh/about",
    )

    # 13. Sensitivity
    s = add(TABLE)
    title(s, 11, "What would change the answer", width=11.01)
    drop(s, 10)
    table(
        s,
        1.16,
        1.6,
        [5.41, 1.4, 1.4, 1.4, 1.4],
        [
            ["Scenario", "Northern", "Upper East", "Upper West", "Nets moved"],
            ["As proposed", "24,196", "18,631", "7,173", "baseline"],
            [
                "Point estimate instead of the upper bound",
                "25,095",
                "17,240",
                "7,665",
                "1,391",
            ],
            [
                "Region-effects model (fits better: AIC 1,253.8 vs 1,286.8)",
                "18,634",
                "16,473",
                "14,893",
                "7,720",
            ],
            [
                "Coverage corrected for the 11 districts now in Savannah or North East",
                "27,686",
                "14,891",
                "7,423",
                "3,740",
            ],
            [
                "Northern coverage at its lower interval bound (61.4%)",
                "22,939",
                "18,596",
                "8,465",
                "1,292",
            ],
            [
                "Northern coverage at its upper interval bound (74.0%)",
                "29,004",
                "11,951",
                "9,045",
                "6,680",
            ],
        ],
        [0.45] + [0.55] * 6,
        size=13,
        right_cols=(1, 2, 3, 4),
        bold_rows=(1,),
    )
    text(
        s,
        1.16,
        5.5,
        11.01,
        0.85,
        [
            P(
                B("Takeaway: ", c=RED),
                "within a region, nets follow population in every scenario except the coverage correction. The "
                "regional split is the uncertain part, so treat regional totals as a range, not a point.",
            )
        ],
        size=15,
    )
    notes(
        s,
        "Say: This is where we volunteer our weak points. A region-effects model, which fits the data "
        "better, moves 7,720 nets. Correcting the coverage of 11 districts moves 3,740. Lowering Northern "
        "Region's coverage gives it fewer nets, not more, because of the positive coverage coefficient.\n"
        "Sources: notebooks/04_allocation.ipynb cell alloc_cd13, which also prints the coverage it uses "
        "(Savannah 79.1%, North East 62.8%; published 79.1% and 62.6%, GDHS 2022 Table 12.1) and the "
        "Northern cluster-bootstrap bounds 61.4% and 74.0% (as in notebooks/02 cell cd39).",
    )

    # 14. Equity
    s = add(TEXT_HEAVY)
    title(s, 11, "Equity: who this helps, and who it misses", size=30, width=W)
    drop(s, 12, 13)
    for i, (head, num, body) in enumerate(
        [
            (
                "Outside the map",
                "207 of 260",
                "of Ghana's districts have no district case data in our package. This allocation says nothing "
                "about them and should not be stretched to them.",
            ),
            (
                "Testing access",
                "65%",
                "median test positivity in Northern Region, the highest of the three, yet it reports the fewest "
                "cases per person (0.27 a year, against 0.88 and 0.80). That points to under-testing, which a "
                "case-based rule would punish.",
            ),
            (
                "Interventions we left out",
                "32 and 24",
                "of the 50 districts had indoor spraying and seasonal chemoprevention in 2014-17, including every "
                "Upper East and Upper West district. Both change burden, and sprayed districts are excluded from "
                "national net campaigns.",
            ),
            (
                "Privacy by design",
                "0",
                "data files or notebook outputs in the project's git history. Licensed DHS household records stay "
                "out of git and out of this deck; everything shown is an aggregate.",
            ),
        ]
    ):
        card(
            s,
            0.62 + (i % 2) * 6.15,
            1.55 + (i // 2) * 2.4,
            5.94,
            2.25,
            [
                P(B(head, c=MUTED, s=14), after=0),
                P(B(num, c=RED, s=30), after=2),
                P(body, s=13),
            ],
        )
    notes(
        s,
        "Say: Equity is about who the rule helps and who it cannot see.\n"
        "Sources: 207 of 260 districts: scripts/verify_claims.py (map join). Test positivity and cases per person "
        "by region: scripts/verify_claims.py section E. Spraying and chemoprevention counts: scripts/verify_claims.py "
        "section E, from data/raw/northern-ghana-districts-routine-data-2014-17.xlsx (these columns are "
        "not in ghana_district_cases.csv). Campaign rules: PMI Ghana Malaria Operational Plan FY2017. "
        "Privacy: git history holds no data files and no notebook outputs (checked 2026-09-22).",
    )

    # 15. Limitations and next steps
    s = add(TEXT_HEAVY)
    title(s, 11, "Limitations, and what we would do next", width=W)
    drop(s, 12, 13)
    card(
        s,
        0.62,
        1.6,
        5.94,
        4.1,
        [
            P(B("Limitations", c=RED, s=18), after=8),
            P(
                "11 districts coded Northern are now in Savannah (79.1% of households own a net) or North East "
                "(62.8%). Correcting this moves 3,740 nets.",
                bullet=True,
                after=8,
            ),
            P(
                "39 of 50 districts report more than one confirmed case per resident over 2014-17: repeat "
                "episodes, care-seeking across district lines, or undercounted populations.",
                bullet=True,
                after=8,
            ),
            P("Coverage is from 2022 and cases from 2014-17.", bullet=True, after=8),
            P(
                "The survey has no district identifier or parasitaemia data, and published ITN intervals are empty.",
                bullet=True,
            ),
        ],
        size=16,
    )
    card(
        s,
        6.77,
        1.6,
        5.94,
        4.1,
        [
            P(B("Before any real distribution", c=RED, s=18), after=8),
            P(
                "Refit with today's regions and district-level spraying and chemoprevention.",
                bullet=True,
                after=10,
            ),
            P(
                "Model risk with region effects, and use coverage only in the gap term.",
                bullet=True,
                after=10,
            ),
            P(
                "Validate by holding out whole regions, and report the spread across seeds.",
                bullet=True,
                after=10,
            ),
            P(
                "Deliver through regional medical stores and CHPS compounds, then survey net use after six months.",
                bullet=True,
            ),
        ],
        size=16,
    )
    notes(
        s,
        "Say: These are the limits we know about, and what we would do before any real distribution.\n"
        "Sources: 11 districts, 3,740 nets, 79.1% and 62.8%: notebooks/04_allocation.ipynb alloc_cd13. "
        "39 of 50 districts: scripts/verify_claims.py section E. "
        "Empty ITN intervals: notebooks/02 cell cd37. No district identifier or parasitaemia: "
        "notebooks/02 cell cd06.",
    )

    # 16. Close
    s = add(CLOSING)
    title(s, 0, "Thank you. Questions?", size=60)
    text(
        s,
        1.67,
        5.7,
        10.0,
        0.5,
        [P("ICS553 Machine Learning Essentials · Group 3 · Ashesi University", c=GOLD)],
        size=16,
        align=PP_ALIGN.CENTER,
    )
    text(
        s,
        1.17,
        6.35,
        11.0,
        0.7,
        [
            P(
                B("AI declaration: ", c=GOLD),
                "we used AI tools to help aggregate our information and to generate this "
                "presentation from our aggregated information. How AI was used is logged in "
                "the project's WORKLOG.md.",
            )
        ],
        size=12,
        color=WHITE,
        align=PP_ALIGN.CENTER,
    )
    notes(
        s,
        "Say: Thank you. Every number in this deck has its source in the speaker notes; we are happy to "
        "take questions on any of them.\n"
        "The footnote is the group's AI declaration. Each member's own AI-use declaration is part of their "
        "individual reflection.",
    )

    # remove the template's nine example slides; their parts are dropped on save
    ids = prs.slides._sldIdLst
    for sld in list(ids)[:n_template]:
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)
    # the template was saved in Slide Master view; open the deck in Normal view instead
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    view = prs.part.part_related_by(RT.VIEW_PROPS)
    view._blob = view.blob.replace(b'lastView="sldMasterView"', b'lastView="sldView"')
    prs.core_properties.title = "Where should 50,000 bed nets go?"
    prs.core_properties.author = "ICS553 Group 3, Ashesi University"
    prs.save(OUT)
    print(f"wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
